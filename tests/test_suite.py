"""
Automated test suite for AI Invoice Studio.
Tests invoice calculations, numbering sequence, template generation,
validation, and PDF export.
"""

import os
import openpyxl
import pytest
from services import database
from services.excel_service import ExcelService
from services.pdf_service import PDFService
from services.invoice_service import InvoiceService
from services.validation import (
    calculate_invoice_totals,
    normalize_date,
    validate_pan,
    validate_ifsc,
    validate_invoice_payload,
)

TEST_DB_PATH = "data/test_studio.db"
TEST_OUTPUT_DIR = "generated_invoices/test_output"


@pytest.fixture(autouse=True)
def setup_teardown_db():
    """Setup clean test database before each test."""
    if os.path.exists(TEST_DB_PATH):
        try:
            os.remove(TEST_DB_PATH)
        except OSError:
            pass

    os.makedirs(TEST_OUTPUT_DIR, exist_ok=True)
    database.init_db(TEST_DB_PATH)
    yield
    # Cleanup test files
    if os.path.exists(TEST_DB_PATH):
        try:
            os.remove(TEST_DB_PATH)
        except OSError:
            pass


def test_calculate_invoice_totals_standard():
    """Test standard single item calculation with discount."""
    res = calculate_invoice_totals(unit_price=2500.0, quantity=1, discount=200.0)
    assert res["quantity"] == 1
    assert res["unit_price"] == 2500.0
    assert res["taxable_value"] == 2500.0
    assert res["discount"] == 200.0
    assert res["subtotal_less_discount"] == 2300.0
    assert res["balance_due"] == 2300.0


def test_calculate_invoice_totals_multiple_quantity():
    """Test multiple quantities calculation."""
    res = calculate_invoice_totals(unit_price=1500.0, quantity=3, discount=500.0)
    assert res["taxable_value"] == 4500.0
    assert res["discount"] == 500.0
    assert res["subtotal_less_discount"] == 4000.0
    assert res["balance_due"] == 4000.0


def test_sequential_invoice_numbering():
    """Test sequential invoice numbers increment properly."""
    no1 = database.get_next_invoice_number(prefix="DL-", db_path=TEST_DB_PATH)
    assert no1 == "DL-001"

    # Insert an invoice
    creators = database.get_creators(db_path=TEST_DB_PATH)
    creator = creators[0]

    database.create_invoice(
        {
            "invoice_number": "DL-001",
            "creator_id": creator["id"],
            "creator_name": creator["name"],
            "invoice_date": "29-10-2026",
            "service_description": "Promotional Video",
            "quantity": 1,
            "unit_price": 2500.0,
            "taxable_value": 2500.0,
            "discount": 200.0,
            "subtotal_less_discount": 2300.0,
            "balance_due": 2300.0,
        },
        db_path=TEST_DB_PATH,
    )

    no2 = database.get_next_invoice_number(prefix="DL-", db_path=TEST_DB_PATH)
    assert no2 == "DL-002"


def test_prevent_duplicate_invoice_number():
    """Test database rejects duplicate invoice numbers."""
    creators = database.get_creators(db_path=TEST_DB_PATH)
    creator = creators[0]

    inv_data = {
        "invoice_number": "DL-100",
        "creator_id": creator["id"],
        "creator_name": creator["name"],
        "invoice_date": "29-10-2026",
        "service_description": "Test Video",
        "quantity": 1,
        "unit_price": 1000.0,
        "taxable_value": 1000.0,
        "discount": 0.0,
        "subtotal_less_discount": 1000.0,
        "balance_due": 1000.0,
    }

    database.create_invoice(inv_data, db_path=TEST_DB_PATH)

    with pytest.raises(ValueError) as excinfo:
        database.create_invoice(inv_data, db_path=TEST_DB_PATH)
    assert "already exists" in str(excinfo.value)


def test_date_normalization():
    """Test natural language dates convert to DD-MM-YYYY."""
    assert normalize_date("29 October 2026") == "29-10-2026"
    assert normalize_date("2026-10-29") == "29-10-2026"
    assert normalize_date("29/10/2026") == "29-10-2026"
    assert normalize_date("28-02-2026") == "28-02-2026"


def test_pan_and_ifsc_validation():
    """Test PAN and IFSC code formats."""
    assert validate_pan("ABCDE1234F") is True
    assert validate_pan("abcde1234f") is True
    assert validate_pan("INVALIDPAN") is False

    assert validate_ifsc("HDFC0000123") is True
    assert validate_ifsc("UTIB0000054") is True
    assert validate_ifsc("1234") is False


def test_excel_generation_matches_template():
    """Test that generated Excel preserves cells and updates values correctly."""
    excel_svc = ExcelService(output_dir=TEST_OUTPUT_DIR)
    assert excel_svc.verify_template() is True

    invoice_data = {
        "invoice_number": "DL-TEST-001",
        "invoice_date": "29-10-2026",
        "service_description": "Promotional Video for Dermalist Skincare",
        "quantity": 1,
        "unit_price": 2500.0,
        "taxable_value": 2500.0,
        "discount": 200.0,
        "video_link": "https://instagram.com/reel/123",
    }
    creator_data = {
        "name": "Aaditi Abhijeet Raut",
        "address": "Sunshine Apts, Mumbai 400053",
        "email": "aaditi@example.com",
        "phone": "+91 98201 23456",
        "pan": "ABCDE1234F",
        "upi_id": "aaditi@okhdfcbank",
        "bank_account": "50100234567890",
        "ifsc": "HDFC0000123",
        "account_holder": "Aaditi Abhijeet Raut",
        "bank_name": "HDFC Bank",
    }

    excel_file = excel_svc.generate_invoice_excel(invoice_data, creator_data)
    assert os.path.exists(excel_file)

    # Open generated file and inspect content
    wb = openpyxl.load_workbook(excel_file, data_only=False)
    ws = wb.active

    assert ws["B2"].value == "Aaditi Abhijeet Raut"
    assert ws["B3"].value == "Sunshine Apts, Mumbai 400053"
    assert ws["D9"].value == "Invoice No: DL-TEST-001"
    assert ws["D10"].value == "Invoice Date:29-10-2026"
    assert ws["B18"].value == "Promotional Video for Dermalist Skincare"
    assert ws["D18"].value == 2500.0
    assert ws["D30"].value == 200.0
    assert ws["D29"].value == "=SUM(D18:D28)"
    assert ws["D31"].value == "=D29-D30"
    # Ensure our formula fix is in place (=D31 instead of double subtraction =D31-D30)
    assert ws["D32"].value == "=D31"
    assert "aaditi@okhdfcbank" in str(ws["B32"].value)
    assert "HDFC0000123" in str(ws["B34"].value)

    wb.close()


def test_pdf_generation_fallback():
    """Test that PDF is generated without error."""
    pdf_svc = PDFService(output_dir=TEST_OUTPUT_DIR)
    invoice_data = {
        "invoice_number": "DL-TEST-001",
        "invoice_date": "29-10-2026",
        "service_description": "Promotional Video",
        "quantity": 1,
        "unit_price": 2500.0,
        "taxable_value": 2500.0,
        "discount": 200.0,
        "subtotal_less_discount": 2300.0,
        "balance_due": 2300.0,
    }
    creator_data = {
        "name": "Aaditi Abhijeet Raut",
        "address": "Mumbai",
        "email": "aaditi@example.com",
        "phone": "9820123456",
        "pan": "ABCDE1234F",
        "upi_id": "aaditi@upi",
        "bank_account": "123456",
        "ifsc": "HDFC0000123",
        "bank_name": "HDFC",
    }

    pdf_path = pdf_svc.generate_pdf(invoice_data, creator_data)
    assert os.path.exists(pdf_path)
    assert os.path.getsize(pdf_path) > 1000


def test_sanitize_row_decimals():
    """Test that _sanitize_row converts Decimal objects to float."""
    from decimal import Decimal
    raw = {
        "invoice_number": "DL-001",
        "balance_due": Decimal("2500.50"),
        "unit_price": Decimal("1000.00"),
        "quantity": 1,
    }
    clean = database._sanitize_row(raw)
    assert isinstance(clean["balance_due"], float)
    assert clean["balance_due"] == 2500.50
    assert isinstance(clean["unit_price"], float)
    assert clean["quantity"] == 1


def test_delete_creator_with_invoices_fails():
    """Test that deleting a creator with existing invoices raises ValueError."""
    creators = database.get_creators(db_path=TEST_DB_PATH)
    creator = creators[0]

    # Add an invoice for this creator
    database.create_invoice(
        {
            "invoice_number": "DL-999",
            "creator_id": creator["id"],
            "creator_name": creator["name"],
            "invoice_date": "29-10-2026",
            "service_description": "Audit Video",
            "quantity": 1,
            "unit_price": 5000.0,
            "taxable_value": 5000.0,
            "discount": 0.0,
            "subtotal_less_discount": 5000.0,
            "balance_due": 5000.0,
        },
        db_path=TEST_DB_PATH,
    )

    with pytest.raises(ValueError) as exc:
        database.delete_creator(creator["id"], db_path=TEST_DB_PATH)
    assert "associated invoices" in str(exc.value)


def test_unique_creator_invoice_generation():
    """Verify that generating an invoice for a new unique creator writes their unique details, not default/Aaditi."""
    excel_svc = ExcelService(output_dir=TEST_OUTPUT_DIR)
    pdf_svc = PDFService(output_dir=TEST_OUTPUT_DIR)
    inv_svc = InvoiceService(excel_service=excel_svc, pdf_service=pdf_svc, db_path=TEST_DB_PATH)

    unique_invoice_data = {
        "invoice_number": "DL-UNIQUE-007",
        "creator_name": "Kunal Kapoor",
        "invoice_date": "30-10-2026",
        "service_description": "Custom Skincare Ad",
        "quantity": 1,
        "unit_price": 7500.0,
        "taxable_value": 7500.0,
        "discount": 500.0,
        "subtotal_less_discount": 7000.0,
        "balance_due": 7000.0,
        "status": "Generated",
    }
    unique_creator_data = {
        "name": "Kunal Kapoor",
        "address": "902 Silver Oak, Bandra, Mumbai 400050",
        "email": "kunal@kapoor.com",
        "phone": "+91 99999 11111",
        "pan": "KUNAL1234K",
        "upi_id": "kunal@kotak",
        "bank_account": "987654321012",
        "ifsc": "KKBK0000123",
        "account_holder": "Kunal Kapoor",
        "bank_name": "Kotak Mahindra Bank",
    }

    result = inv_svc.generate_and_save_invoice(unique_invoice_data, unique_creator_data)

    # Check Excel file content
    wb = openpyxl.load_workbook(result["excel_path"], data_only=False)
    ws = wb.active

    assert ws["B2"].value == "Kunal Kapoor"
    assert "Bandra" in ws["B3"].value
    assert "KUNAL1234K" in ws["B7"].value
    assert "kunal@kotak" in ws["B32"].value
    assert "987654321012" in ws["B33"].value
    assert "KKBK0000123" in ws["B34"].value
    assert "Aaditi" not in str(ws["B2"].value)
    wb.close()

