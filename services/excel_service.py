"""
Excel template manipulation service for AI Invoice Studio.
Fills the master Excel template while strictly preserving fonts, borders,
merged cells, formatting, and dimensions. Never overwrites the master template.
"""

import os
import shutil
from typing import Any, Dict
import openpyxl
from dotenv import load_dotenv

load_dotenv()

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

raw_tpl = os.getenv("TEMPLATE_PATH", "templates/non gst invoice format - Dermalist (2).xlsx")
if not os.path.isabs(raw_tpl):
    DEFAULT_TEMPLATE_PATH = os.path.join(BASE_DIR, raw_tpl)
else:
    DEFAULT_TEMPLATE_PATH = raw_tpl

raw_out = os.getenv("OUTPUT_DIR", "generated_invoices")
if not os.path.isabs(raw_out):
    DEFAULT_OUTPUT_DIR = os.path.join(BASE_DIR, raw_out)
else:
    DEFAULT_OUTPUT_DIR = raw_out


class ExcelService:
    def __init__(self, template_path: str = DEFAULT_TEMPLATE_PATH, output_dir: str = DEFAULT_OUTPUT_DIR):
        self.template_path = os.path.abspath(template_path)
        self.output_dir = os.path.abspath(output_dir)
        os.makedirs(self.output_dir, exist_ok=True)

    def verify_template(self) -> bool:
        """Verify that the master template exists and is readable."""
        return os.path.exists(self.template_path) and os.path.isfile(self.template_path)

    def generate_invoice_excel(
        self,
        invoice_data: Dict[str, Any],
        creator_data: Dict[str, Any],
        output_filename: str = None,
    ) -> str:
        """
        Create a new Excel invoice workbook populated from invoice and creator data.
        Returns the absolute filepath of the generated workbook.
        """
        if not self.verify_template():
            raise FileNotFoundError(f"Master Excel template not found at '{self.template_path}'")

        inv_num = invoice_data.get("invoice_number", "INV-001").replace("/", "_").replace("\\", "_")
        if not output_filename:
            output_filename = f"Invoice_{inv_num}.xlsx"

        output_path = os.path.join(self.output_dir, output_filename)

        # Load master workbook (data_only=False preserves formulas)
        wb = openpyxl.load_workbook(self.template_path, data_only=False)
        ws = wb.active

        # 1. Creator Header Information
        creator_name = creator_data.get("name", "")
        creator_address = creator_data.get("address", "")
        creator_email = creator_data.get("email", "")
        creator_phone = creator_data.get("phone", "")
        creator_pan = creator_data.get("pan", "")

        if creator_name:
            ws["B2"].value = creator_name

        if creator_address:
            ws["B3"].value = creator_address

        ws["B5"].value = f"Email ID: {creator_email}" if creator_email else "Email ID:"
        ws["B6"].value = f"Contact Number: {creator_phone}" if creator_phone else "Contact Number"
        ws["B7"].value = f"Pan Number : {creator_pan}" if creator_pan else "Pan Number : "

        # 2. Invoice Metadata
        inv_no_val = invoice_data.get("invoice_number", "")
        inv_date_val = invoice_data.get("invoice_date", "")

        ws["D9"].value = f"Invoice No: {inv_no_val}"
        ws["D10"].value = f"Invoice Date:{inv_date_val}"

        # 3. Line Items & Amounts
        description = invoice_data.get("service_description", "")
        unit_price = float(invoice_data.get("unit_price", 0.0))
        quantity = int(invoice_data.get("quantity", 1))
        taxable_value = float(invoice_data.get("taxable_value", unit_price * quantity))
        discount = float(invoice_data.get("discount", 0.0))
        video_link = invoice_data.get("video_link", "")

        # Service description row
        desc_line = description
        if quantity > 1:
            desc_line = f"{quantity} x {description}"
        ws["B18"].value = desc_line
        ws["D18"].value = taxable_value

        # Video link row (if provided)
        if video_link:
            ws["B20"].value = f"Video Link- {video_link}"
        else:
            ws["B20"].value = "Video Link- "

        # Blank out unused item rows D19-D28 to ensure clean calculation
        for r in range(19, 29):
            if r != 20:  # row 20 has video link
                ws[f"B{r}"].value = None
            ws[f"D{r}"].value = None

        # 4. Formulas and Totals
        # Subtotal formula
        ws["D29"].value = "=SUM(D18:D28)"

        # Discount
        ws["D30"].value = discount

        # Subtotal Less Discount formula
        ws["D31"].value = "=D29-D30"

        # Balance Due formula:
        # Corrected template calculation: original had =D31-D30 (double subtraction bug).
        # Since D31 is already Subtotal less Discount, D32 must be =D31.
        ws["D32"].value = "=D31"

        # 5. Creator Payment & Bank Details
        upi_id = creator_data.get("upi_id", "")
        bank_account = creator_data.get("bank_account", "")
        ifsc = creator_data.get("ifsc", "")
        acc_holder = creator_data.get("account_holder", creator_name)
        bank_name = creator_data.get("bank_name", "")

        ws["B32"].value = f"UPI ID : {upi_id}" if upi_id else "UPI ID :"
        ws["B33"].value = f"Bank Account No:- {bank_account}" if bank_account else "Bank Account No:-"

        ws["B34"].value = (
            f"IFSC:- {ifsc}                               "
            f"Name of Accountholder:- {acc_holder}                                    "
            f"Name of Bank:- {bank_name}"
        )

        # Save to output file
        wb.save(output_path)
        wb.close()

        return os.path.abspath(output_path)
